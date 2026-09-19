import os
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
import requests
from bs4 import BeautifulSoup
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches

app = Flask(__name__)
# आपके Blogger से आ रही रिक्वेस्ट को अनुमति देने के लिए
CORS(app)

# SlideShare को यह दिखाने के लिए कि रिक्वेस्ट ब्राउज़र से आ रही है
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

@app.route("/")
def home():
    return "Backend status: Running and Active!"

@app.route("/download")
def download():
    url = request.args.get("url")
    if not url:
        return jsonify({"error": "SlideShare URL is missing"}), 400

    if "slideshare.net" not in url:
        return jsonify({"error": "Only valid SlideShare links are supported"}), 400

    try:
        # 1. SlideShare पेज का HTML सोर्स मंगाएं
        response = requests.get(url, headers=HEADERS, timeout=20)
        if response.status_code != 200:
            return jsonify({"error": f"Failed to open page. Status: {response.status_code}"}), 400

        soup = BeautifulSoup(response.text, "html.parser")
        
        # 2. सभी स्लाइड्स की इमेजेस के लिंक्स ढूंढें
        images = soup.find_all("img", class_="SlideImage")
        if not images:
            images = [img for img in soup.find_all("img") if "slide" in img.get("src", "").lower() or img.get("data-full")]

        if not images:
            return jsonify({"error": "Slides could not be extracted from this page"}), 404

        # 3. नया PowerPoint स्ट्रक्चर तैयार करें
        prs = Presentation()
        prs.slide_width = Inches(13.33)  # 16:9 Widescreen Layout
        prs.slide_height = Inches(7.5)
        
        # इंडेक्स [6] का उपयोग करके खाली (Blank) स्लाइड लेआउट चुनें
        blank_layout = prs.slide_layouts[6]

        # 4. हर एक स्लाइड इमेज को डाउनलोड करके PPTX में जोड़ें
        for img_tag in images:
            img_url = img_tag.get("data-full") or img_tag.get("data-normal") or img_tag.get("src")
            if not img_url:
                continue

            # इमेज को सीधे RAM (Memory) में स्टोर करें
            img_data = requests.get(img_url, headers=HEADERS, timeout=10).content
            img_stream = BytesIO(img_data)

            # PPTX में जोड़ें
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(img_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

        # 5. फाइनल फाइल को सीधे यूजर के ब्राउज़र में भेजें
        output_stream = BytesIO()
        prs.save(output_stream)
        output_stream.seek(0)

        return send_file(
            output_stream,
            mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            as_attachment=True,
            download_name="slideshare_presentation.pptx"
        )

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    # Render पोर्ट बाइंडिंग
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
